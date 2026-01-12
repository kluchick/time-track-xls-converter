"""
Excel Time Tracker Converter

Reads time tracking data from input.xlsx and writes converted data
to a dated copy of output.xlsx on the Efforts sheet.
"""

import os
import shutil
from datetime import datetime
from pathlib import Path

import pandas as pd
import openpyxl
from openpyxl.utils import get_column_letter


# File paths
BASE_DIR = Path(__file__).parent.parent
FILES_DIR = BASE_DIR / "files"
INPUT_FILE = FILES_DIR / "input.xlsx"
OUTPUT_TEMPLATE = FILES_DIR / "output.xlsx"

# Column mappings
INPUT_COLUMNS = ["Project Task", "Date", "Duration"]
OUTPUT_COLUMNS = ["Effort", "Description", "Date"]
EFFORTS_SHEET_NAME = "Efforts"


def validate_input_file():
    """Check if input file exists and has required columns."""
    if not INPUT_FILE.exists():
        raise FileNotFoundError(f"Input file not found: {INPUT_FILE}")
    
    try:
        df = pd.read_excel(INPUT_FILE, nrows=0)  # Read only headers
        missing_cols = [col for col in INPUT_COLUMNS if col not in df.columns]
        if missing_cols:
            raise ValueError(f"Missing required columns in input file: {missing_cols}")
    except Exception as e:
        raise ValueError(f"Error reading input file: {e}")


def create_dated_output_file():
    """Create a copy of output.xlsx with current date in filename."""
    if not OUTPUT_TEMPLATE.exists():
        raise FileNotFoundError(f"Output template not found: {OUTPUT_TEMPLATE}")
    
    date_str = datetime.now().strftime("%Y-%m-%d")
    output_filename = f"output_{date_str}.xlsx"
    output_path = FILES_DIR / output_filename
    
    # Copy template file
    shutil.copy2(OUTPUT_TEMPLATE, output_path)
    print(f"Created dated output file: {output_path}")
    
    return output_path


def read_input_data():
    """Read and validate input data."""
    df = pd.read_excel(INPUT_FILE)
    
    # Check for required columns
    missing_cols = [col for col in INPUT_COLUMNS if col not in df.columns]
    if missing_cols:
        raise ValueError(f"Missing required columns: {missing_cols}")
    
    # Select only required columns
    df = df[INPUT_COLUMNS].copy()
    
    # Validate Duration is numeric
    df["Duration"] = pd.to_numeric(df["Duration"], errors="coerce")
    invalid_duration = df["Duration"].isna()
    if invalid_duration.any():
        print(f"Warning: Found {invalid_duration.sum()} rows with invalid Duration values. They will be skipped.")
        df = df[~invalid_duration].copy()
    
    # Remove rows with empty required fields
    df = df.dropna(subset=["Project Task", "Date", "Duration"])
    
    print(f"Read {len(df)} valid rows from input file")
    return df


def process_data(df):
    """Convert input data to output format."""
    processed = pd.DataFrame()
    
    # Convert Duration (minutes) to Effort (hours)
    processed["Effort"] = df["Duration"] / 60.0
    
    # Map Project Task to Description
    processed["Description"] = df["Project Task"]
    
    # Keep Date as-is
    processed["Date"] = df["Date"]
    
    return processed


def write_to_efforts_sheet(output_path, data_df):
    """Write processed data to Efforts sheet in output file."""
    # Open workbook
    wb = openpyxl.load_workbook(output_path)
    
    # Get or create Efforts sheet
    if EFFORTS_SHEET_NAME not in wb.sheetnames:
        ws = wb.create_sheet(EFFORTS_SHEET_NAME)
        print(f"Created '{EFFORTS_SHEET_NAME}' sheet")
    else:
        ws = wb[EFFORTS_SHEET_NAME]
    
    # Data starts from row 3 (row 1 is template header, row 2 is column headers)
    data_start_row = 3
    
    # Clear existing data in columns B, C, D only (preserve column A)
    if ws.max_row >= data_start_row:
        for row in range(data_start_row, ws.max_row + 1):
            for col in range(2, 5):  # Columns B, C, D (indices 2, 3, 4)
                ws.cell(row, col).value = None
    
    # Write data starting from row 3, columns B, C, D
    # Map: Effort -> B (index 2), Description -> C (index 3), Date -> D (index 4)
    for row_idx, (_, row_data) in enumerate(data_df.iterrows(), start=data_start_row):
        for col_idx, col_name in enumerate(OUTPUT_COLUMNS, start=2):  # start=2 for column B
            value = row_data[col_name]
            
            # Handle date formatting
            if col_name == "Date" and isinstance(value, pd.Timestamp):
                ws.cell(row_idx, col_idx).value = value
                ws.cell(row_idx, col_idx).number_format = "YYYY-MM-DD"
            else:
                ws.cell(row_idx, col_idx).value = value
    
    # Save workbook
    wb.save(output_path)
    print(f"Written {len(data_df)} rows to '{EFFORTS_SHEET_NAME}' sheet (columns B, C, D)")


def main():
    """Main conversion function."""
    try:
        print("Starting Excel Time Tracker Converter...")
        
        # Validate input file
        validate_input_file()
        
        # Read input data
        input_df = read_input_data()
        
        if input_df.empty:
            print("Warning: No valid data found in input file. Creating empty output file.")
        
        # Process data
        processed_df = process_data(input_df)
        
        # Create dated output file
        output_path = create_dated_output_file()
        
        # Write to Efforts sheet
        write_to_efforts_sheet(output_path, processed_df)
        
        print(f"\nConversion completed successfully!")
        print(f"Output file: {output_path}")
        
    except Exception as e:
        print(f"\nError: {e}")
        raise


if __name__ == "__main__":
    main()

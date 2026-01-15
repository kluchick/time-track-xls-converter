"""
Excel Time Tracker Converter - Core Business Logic

Refactored core functionality for converting time tracking Excel files.
Separated from UI layer for better modularity and testability.
"""

from pathlib import Path
from datetime import datetime
from typing import Tuple
from io import BytesIO

import pandas as pd
import openpyxl


# Column mappings
INPUT_COLUMNS = ["Project Task", "Date", "Duration"]
OUTPUT_COLUMNS = ["Effort", "Description", "Date"]
EFFORTS_SHEET_NAME = "Efforts"


class ConversionError(Exception):
    """Custom exception for conversion errors"""
    pass


def validate_input_data(file_path: Path) -> Tuple[bool, str]:
    """
    Validate input file exists and has required columns.

    Args:
        file_path: Path to input Excel file

    Returns:
        Tuple of (is_valid, error_message)
        - (True, "") if valid
        - (False, "error message") if invalid
    """
    # Check if file exists
    if not file_path.exists():
        return False, f"File not found: {file_path}"

    # Check if it's a file (not directory)
    if not file_path.is_file():
        return False, f"Path is not a file: {file_path}"

    # Check file extension
    if file_path.suffix.lower() not in ['.xlsx', '.xls']:
        return False, "File must be an Excel file (.xlsx or .xls)"

    # Try to read and validate columns
    try:
        df = pd.read_excel(file_path, nrows=0)  # Read only headers
        missing_cols = [col for col in INPUT_COLUMNS if col not in df.columns]
        if missing_cols:
            return False, f"Missing required columns: {', '.join(missing_cols)}"
        return True, ""
    except Exception as e:
        return False, f"Error reading file: {str(e)}"


def read_and_process_input(file_path: Path, extend_description: str = None) -> pd.DataFrame:
    """
    Read input Excel and process data.

    Args:
        file_path: Path to input Excel file
        extend_description: Optional description text for extended time entries.
                           If provided and not empty, adds filler rows for days with < 8 hours.

    Returns:
        Processed DataFrame with columns: Effort, Description, Date

    Raises:
        ConversionError: If reading or processing fails
    """
    try:
        # Read input file
        df = pd.read_excel(file_path)

        # Check for required columns
        missing_cols = [col for col in INPUT_COLUMNS if col not in df.columns]
        if missing_cols:
            raise ConversionError(f"Missing required columns: {', '.join(missing_cols)}")

        # Select only required columns
        df = df[INPUT_COLUMNS].copy()

        # Validate Duration is numeric
        df["Duration"] = pd.to_numeric(df["Duration"], errors="coerce")
        invalid_duration = df["Duration"].isna()
        if invalid_duration.any():
            print(f"Warning: Found {invalid_duration.sum()} rows with invalid Duration values. They will be skipped.")
            df = df[~invalid_duration].copy()

        # Remove rows with empty required fields
        df = df.dropna(subset=["Project Task", "Date", "Duration"])

        if df.empty:
            raise ConversionError("No valid data found in input file")

        # Create output DataFrame
        processed = pd.DataFrame()

        # Convert Duration (minutes) to Effort (hours)
        processed["Effort"] = df["Duration"] / 60.0

        # Map Project Task to Description
        processed["Description"] = df["Project Task"]

        # Keep Date as-is
        processed["Date"] = df["Date"]

        # Normalize dates for grouping (remove time component)
        processed["Date"] = pd.to_datetime(processed["Date"]).dt.normalize()

        # Extend time entries if description is provided
        if extend_description and extend_description.strip():
            processed = extend_time_entries(processed, extend_description.strip())

        print(f"Processed {len(processed)} rows from input file")
        return processed

    except ConversionError:
        raise
    except Exception as e:
        raise ConversionError(f"Error processing input file: {str(e)}")


def extend_time_entries(processed_df: pd.DataFrame, extend_description: str) -> pd.DataFrame:
    """
    Extend time entries to ensure each day has at least 8 hours.

    Args:
        processed_df: DataFrame with columns: Effort, Description, Date
        extend_description: Description text for filler rows

    Returns:
        Extended DataFrame with filler rows added for days with < 8 hours
    """
    # Sort by date first to maintain chronological order
    processed_df = processed_df.sort_values("Date").reset_index(drop=True)

    # Group by date and calculate total hours per day
    daily_totals = processed_df.groupby("Date")["Effort"].sum()

    # Find days that need extension (< 8 hours)
    days_to_extend = daily_totals[daily_totals < 8.0]

    if len(days_to_extend) == 0:
        return processed_df

    # Build extended DataFrame by inserting filler rows after each day's entries
    extended_rows = []
    current_date = None
    
    for idx, row in processed_df.iterrows():
        row_date = row["Date"]
        
        # If we've moved to a new day and the previous day needed extension
        if current_date is not None and current_date != row_date and current_date in days_to_extend.index:
            # Add filler row for the previous day
            day_total = daily_totals[current_date]
            missing_hours = 8.0 - day_total
            extended_rows.append({
                "Effort": missing_hours,
                "Description": extend_description,
                "Date": current_date
            })
        
        extended_rows.append(row.to_dict())
        current_date = row_date
    
    # Handle the last day if it needs extension
    if current_date is not None and current_date in days_to_extend.index:
        day_total = daily_totals[current_date]
        missing_hours = 8.0 - day_total
        extended_rows.append({
            "Effort": missing_hours,
            "Description": extend_description,
            "Date": current_date
        })

    # Create DataFrame from extended rows
    extended_df = pd.DataFrame(extended_rows)

    print(f"Extended {len(days_to_extend)} days to 8 hours with description: '{extend_description}'")
    return extended_df


def write_to_output_file(template_bytes: bytes, data_df: pd.DataFrame, output_path: Path):
    """
    Load template from bytes, write data to Efforts sheet, save to output_path.

    Args:
        template_bytes: Bytes containing the Excel template
        data_df: DataFrame with columns: Effort, Description, Date
        output_path: Path where to save the output file

    Raises:
        ConversionError: If writing fails
    """
    try:
        # Load workbook from bytes
        template_stream = BytesIO(template_bytes)
        wb = openpyxl.load_workbook(template_stream)

        # Get or create Efforts sheet
        if EFFORTS_SHEET_NAME not in wb.sheetnames:
            ws = wb.create_sheet(EFFORTS_SHEET_NAME)
            print(f"Created '{EFFORTS_SHEET_NAME}' sheet")
        else:
            ws = wb[EFFORTS_SHEET_NAME]

        # Data starts from row 3 (row 1 is template header, row 2 is column headers)
        data_start_row = 3

        # Find the last row with data in columns B, C, or D
        last_data_row = data_start_row - 1  # Start from row 2 (before data starts)
        for row in range(data_start_row, ws.max_row + 1):
            # Check if any of columns B, C, D have data
            has_data = any(
                ws.cell(row, col).value is not None and str(ws.cell(row, col).value).strip() != ""
                for col in range(2, 5)  # Columns B, C, D
            )
            if has_data:
                last_data_row = row

        # Start writing new data after the last existing data row
        write_start_row = last_data_row + 1

        # Check if column A in first data row has a value to copy to all imported rows
        column_a_value = ws.cell(data_start_row, 1).value  # A3
        copy_column_a = column_a_value is not None and str(column_a_value).strip() != ""

        # Write data starting after last existing row, columns B, C, D
        # Map: Effort -> B (index 2), Description -> C (index 3), Date -> D (index 4)
        for row_idx, (_, row_data) in enumerate(data_df.iterrows(), start=write_start_row):
            # If column A has a value in first data row, copy it to all imported rows
            if copy_column_a:
                ws.cell(row_idx, 1).value = column_a_value  # Column A

            for col_idx, col_name in enumerate(OUTPUT_COLUMNS, start=2):  # start=2 for column B
                value = row_data[col_name]

                # Handle date formatting
                if col_name == "Date" and isinstance(value, pd.Timestamp):
                    ws.cell(row_idx, col_idx).value = value
                    ws.cell(row_idx, col_idx).number_format = "YYYY-MM-DD"
                else:
                    ws.cell(row_idx, col_idx).value = value

        # Save workbook
        wb.save(output_path)
        if copy_column_a:
            print(f"Appended {len(data_df)} rows to '{EFFORTS_SHEET_NAME}' sheet starting from row {write_start_row} (columns A, B, C, D)")
            print(f"  - Copied column A value '{column_a_value}' to all imported rows")
        else:
            print(f"Appended {len(data_df)} rows to '{EFFORTS_SHEET_NAME}' sheet starting from row {write_start_row} (columns B, C, D)")

    except Exception as e:
        raise ConversionError(f"Error writing output file: {str(e)}")


def convert_excel_file(input_path: Path, template_bytes: bytes, output_dir: Path, extend_description: str = None) -> Path:
    """
    Main conversion function. Converts input Excel to output with timestamp.

    Args:
        input_path: Path to input Excel file
        template_bytes: Bytes containing the Excel template
        output_dir: Directory where to save the output file
        extend_description: Optional description text for extended time entries.
                           If provided and not empty, adds filler rows for days with < 8 hours.

    Returns:
        Path to the created output file

    Raises:
        ConversionError: If conversion fails
    """
    try:
        # Validate input
        is_valid, error_msg = validate_input_data(input_path)
        if not is_valid:
            raise ConversionError(error_msg)

        # Read and process data
        processed_df = read_and_process_input(input_path, extend_description=extend_description)

        # Generate timestamped output filename
        timestamp = datetime.now().strftime("%Y-%m-%d_%H-%M-%S")
        output_filename = f"output_{timestamp}.xlsx"
        output_path = output_dir / output_filename

        # Ensure output directory exists
        output_dir.mkdir(parents=True, exist_ok=True)

        # Write to output file
        write_to_output_file(template_bytes, processed_df, output_path)

        print(f"Conversion completed successfully!")
        print(f"Output file: {output_path}")

        return output_path

    except ConversionError:
        raise
    except Exception as e:
        raise ConversionError(f"Unexpected error during conversion: {str(e)}")
